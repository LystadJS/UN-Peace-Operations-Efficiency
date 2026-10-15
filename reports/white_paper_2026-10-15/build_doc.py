"""Author a paginated research white paper from reviewed Markdown and figures."""
from pathlib import Path
import re,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).parent;OUT=ROOT/'deliverables';OUT.mkdir(exist_ok=True)
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.67);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.73);sec.right_margin=Inches(.73);sec.header_distance=Inches(.25);sec.footer_distance=Inches(.27);sec.different_first_page_header_footer=True
styles=doc.styles
for name in ['Normal','Body Text']:
 st=styles[name];st.font.name='Liberation Sans';st.font.size=Pt(10.5);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.10
for name,size in [('Title',38),('Subtitle',20),('Heading 1',21),('Heading 2',14),('Heading 3',11)]:
 st=styles[name];st.font.name='Liberation Sans';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('153C52');st.paragraph_format.space_after=Pt(10);st.paragraph_format.space_before=Pt(0);st.paragraph_format.keep_with_next=True
st=styles['Caption'];st.font.name='Liberation Sans';st.font.size=Pt(8.6);st.font.color.rgb=RGBColor.from_string('44535B');st.paragraph_format.space_after=Pt(8);st.paragraph_format.line_spacing=1.05
header=sec.header.paragraphs[0];header.text='UN PEACE OPERATIONS EFFICIENCY  /  RESEARCH WHITE PAPER';header.style='Caption';header.paragraph_format.space_after=Pt(0)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('15 October 2026  |  ');r.font.size=Pt(8);r.font.name='Liberation Sans'
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
props=doc.core_properties;props.title='Shared Inputs, Distinct Mandates';props.subject='Portfolio assessment of UN peace operations efficiency';props.author='UN Peace Operations Efficiency Project';props.keywords='Peacekeeping; DPPA; DPO; mandates; financing; research reference 13.0';props.comments='Independent project synthesis; not an official United Nations publication.'
def hyperlink(p,text,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'175A78');pr.append(co);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)
def inline(p,s):
 pattern=r'(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|`[^`]+`)'
 for t in re.split(pattern,s):
  if not t:continue
  if t.startswith('[') and '](' in t:
   m=re.match(r'\[([^\]]+)\]\((https?://.+)\)',t)
   if m:hyperlink(p,m[1],m[2]);continue
  rr=p.add_run(t[2:-2] if t.startswith('**') else t[1:-1] if t.startswith('`') else t)
  if t.startswith('**'):rr.bold=True
  if t.startswith('`'):rr.font.name='Liberation Mono';rr.font.size=Pt(8.5)
def table(lines):
 data=[[x.strip() for x in line.strip().strip('|').split('|')] for line in lines if not re.match(r'^\|[\s:|\-]+\|$',line)]
 n=len(data[0]);tab=doc.add_table(rows=1,cols=n);tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
 if n==2:weights=[.33,.67]
 elif n==3:weights=[.24,.35,.41]
 elif n==4:weights=[.19,.25,.27,.29]
 else:weights=[1/n]*n
 if data[0][0]=='Rank':weights=[.07,.08,.35,.50]
 for rowidx,rowdata in enumerate(data):
  row=tab.rows[0] if rowidx==0 else tab.add_row()
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if rowidx==0:
   rep=OxmlElement('w:tblHeader');rep.set(qn('w:val'),'true');trpr.append(rep)
  for j,txt in enumerate(rowdata):
   c=row.cells[j];c.width=Inches(7.04*weights[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=c.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.03
   inline(p,txt)
   for rr in p.runs:rr.font.size=Pt(9.0 if n>=4 else 9.3)
   tcpr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for k,val in [('top',45),('left',70),('bottom',45),('right',70)]:
    el=OxmlElement('w:'+k);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');mar.append(el)
   tcpr.append(mar)
   if rowidx==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'153C52');tcpr.append(sh)
    for rr in p.runs:rr.bold=True;rr.font.color.rgb=RGBColor(255,255,255)
   elif rowidx%2==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F0F4F6');tcpr.append(sh)
 for col,w in zip(tab.columns,weights):col.width=Inches(7.04*w)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return tab
raw=(ROOT/'whitepaper.md').read_text();pages=raw.split('<!-- page -->')
contents='''# Reading guide

The main text is organized around decisions, not the project's version history. Tables state evidence depth and safeguards; figures show the underlying recorded quantities without treating them as operational outcomes.

| Read for | Sections | Page |
|---|---|---:|
| Whole-project status and limits | 1–3 | 4–6 |
| Peacekeeping and support-office financing | 4–8 | 7–11 |
| SPM proposals and advisory coverage | 9–12 | 12–15 |
| Headquarters and UN80 boundaries | 13–14 | 16–17 |
| All functional families and frozen priorities | 15–17 | 18–20 |
| Census, identity and training evidence | 18–20 | 21–23 |
| Worked pilots and reference quality | 21–24 | 24–27 |
| Owner evidence, shadow protocol and decisions | 25–27 | 28–30 |
| Sources, scope and reproducibility | Evidence notes | 31–32 |

## How to interpret the evidence

**Recorded fact:** a documented mandate, source entry, proposal or appropriation within the stated perimeter.

**Analytical inference:** a reasoned comparison or opportunity that has not been established as actual duplicated work.

**Research prototype:** a representation or software check on selected public-source inputs, not an operational trial.

**Decision still gated:** owner permission, actual work substitution or financial substantiation remains missing.

**Important:** dollars, publications, events, meetings, people and research identifiers are different units. They are not combined into a single efficiency score.

The companion data pack contains the values behind all 19 figures, the complete 59-entry functional mapping and the pinned source-file manifest. The white paper creates no new research-reference version and makes no changes to mandates, rankings or live-use gates.
'''
pages.insert(2,contents)
pageinfo=[]
for idx,page in enumerate(pages):
 if idx:doc.add_page_break()
 lines=page.strip().splitlines();i=0;first=True
 if idx==0:
  doc.add_paragraph().paragraph_format.space_after=Pt(45)
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('|'):
   ls=[]
   while i<len(lines) and lines[i].strip().startswith('|'):ls.append(lines[i]);i+=1
   table(ls);continue
  if line.startswith('!['):
   m=re.match(r'!\[(.+)\]\((.+)\)',line);caption,path=m.groups();p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
   shape=p.add_run().add_picture(str(ROOT/path),width=Inches(6.85));shape._inline.docPr.set('descr',caption)
   cap=doc.add_paragraph(style='Caption');inline(cap,caption);i+=1;continue
  if line.startswith('# '):
   title=line[2:];p=doc.add_paragraph(title,style='Title' if idx==0 else 'Heading 1');pageinfo.append({'planned_page':idx+1,'title':title})
   if idx==0:p.paragraph_format.space_after=Pt(0)
  elif line.startswith('## '):
   p=doc.add_paragraph(line[3:],style='Subtitle' if idx==0 else 'Heading 2')
   if idx==0:p.paragraph_format.space_before=Pt(20);p.paragraph_format.space_after=Pt(26)
  else:
   p=doc.add_paragraph();inline(p,line)
   if idx==0:p.paragraph_format.space_after=Pt(18)
   if idx>=len(pages)-2:
    p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1.04
    for rr in p.runs:rr.font.size=Pt(9.0)
  i+=1
out=OUT/'UN_Peace_Operations_White_Paper_2026-10-15.docx';doc.save(out)
(ROOT/'analysis/planned_pages.json').write_text(json.dumps(pageinfo,indent=2))
print(out,'planned pages',len(pages),'words',len(raw.split()))
